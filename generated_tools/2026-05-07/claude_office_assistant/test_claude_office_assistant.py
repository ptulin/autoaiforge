import pytest
from unittest.mock import patch, MagicMock
from claude_office_assistant import summarize_word, analyze_excel, generate_presentation
import os
from docx import Document
from openpyxl import Workbook

def test_summarize_word(tmp_path):
    # Create a temporary Word document
    file_path = tmp_path / "test.docx"
    document = Document()
    document.add_paragraph("This is a test document.")
    document.save(file_path)

    mock_api_response = MagicMock()
    mock_api_response.choices = [MagicMock(text="This is a summary.")]

    with patch("openai.Completion.create", return_value=mock_api_response):
        result = summarize_word(file_path, "fake_api_key")
        assert result == "This is a summary."

def test_analyze_excel(tmp_path):
    file_path = tmp_path / "test.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["Header1", "Header2"])
    sheet.append([1, 2])
    workbook.save(file_path)

    result = analyze_excel(file_path)
    assert "Sheet1" in result
    assert result["Sheet1"]["row_count"] == 2
    assert result["Sheet1"]["column_count"] == 2

def test_generate_presentation(tmp_path):
    output_file = tmp_path / "output.pptx"
    text = "Title1\nContent1\n\nTitle2\nContent2"

    result = generate_presentation(text, output_file)
    assert os.path.exists(output_file)
    assert result == f"Presentation saved to {output_file}"